"""Texto e notas do orador de um .pptx, no mesmo formato do extrair.sh.
Uso (pelo extrair.sh, com o venv de _shared/.venv): python -I extrair_pptx.py <ficheiro.pptx>
"""
import re
import sys
from collections import Counter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def textos(shapes, tem_imagem):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from textos(sh.shapes, tem_imagem)
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            tem_imagem[0] = True
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                yield " │ ".join(c.text.strip() for c in row.cells)
        elif sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if t:
                    yield ("  " * p.level) + t


prs = Presentation(sys.argv[1])
slides = []
for slide in prs.slides:
    img = [False]
    linhas = list(textos(slide.shapes, img))
    notas = ""
    if slide.has_notes_slide:
        notas = slide.notes_slide.notes_text_frame.text.strip()
    slides.append((linhas, notas, img[0]))

# Rodapés: linhas que se repetem (algarismos à parte) em >=50% dos slides,
# como no extrair.sh. Saem do texto e anunciam-se uma vez no topo.
chave = lambda l: re.sub(r"\d+", "#", l)
cnt = Counter(k for linhas, _, _ in slides for k in {chave(l) for l in linhas if len(l) >= 8})
n = len(slides)
rem = {k for k, c in cnt.items() if n >= 4 and c >= 3 and c >= n / 2}
exemplo = {}
for linhas, _, _ in slides:
    for l in linhas:
        exemplo.setdefault(chave(l), l)
for k in sorted(rem):
    print(f"[rodapé removido: {exemplo[k]}]")

ver = []
for i, (linhas, notas, tem_img) in enumerate(slides, 1):
    linhas = [l for l in linhas if chave(l) not in rem
              and not (re.fullmatch(r"\d{1,3}", l) and l == str(i))]
    img = [tem_img]
    palavras = sum(len(l.split()) for l in linhas)
    marca = " [IMAGEM]" if img[0] else (" [POUCO TEXTO]" if palavras < 15 else "")
    if marca:
        ver.append(str(i))
    print(f"--- slide {i}{marca} ---")
    print("\n".join(linhas))
    if notas:
        print("[notas do orador]\n" + notas)
print()
print(f"=== {len(prs.slides)} páginas. Ver como imagem só se o texto não chegar: {' '.join(ver) or 'nenhuma'}")
print("=== (imagens: extrair com shape.image.blob no venv _shared/.venv; ver só as que importam)")
